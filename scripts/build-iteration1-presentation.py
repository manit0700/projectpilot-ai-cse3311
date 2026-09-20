from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "presentations" / "ProjectPilot_AI_Iteration1_Demo_First.pptx"
ASSETS = ROOT / "docs" / "presentations" / "assets"

WIDE = (13.333, 7.5)

COLORS = {
    "ink": RGBColor(15, 23, 42),
    "muted": RGBColor(71, 85, 105),
    "soft": RGBColor(241, 245, 249),
    "line": RGBColor(203, 213, 225),
    "teal": RGBColor(0, 121, 107),
    "teal_dark": RGBColor(0, 83, 75),
    "teal_soft": RGBColor(224, 251, 247),
    "amber": RGBColor(146, 64, 14),
    "white": RGBColor(255, 255, 255),
}

FONT = "Aptos"


def add_text(slide, text, x, y, w, h, size=22, color="ink", bold=False, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = MSO_ANCHOR.TOP
    p = frame.paragraphs[0]
    p.text = text
    p.font.name = FONT
    p.font.size = Pt(size)
    p.font.color.rgb = COLORS[color]
    p.font.bold = bold
    if align:
        p.alignment = align
    return box


def add_title(slide, title, subtitle=None):
    add_text(slide, title, 0.65, 0.42, 9.7, 0.55, size=28, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.68, 0.97, 9.8, 0.36, size=13, color="muted")


def add_footer(slide, n):
    add_text(slide, f"{n}", 12.45, 7.03, 0.35, 0.22, size=9, color="muted", align=PP_ALIGN.RIGHT)
    add_text(slide, "ProjectPilot AI", 0.65, 7.03, 2.2, 0.22, size=9, color="muted")


def add_rect(slide, x, y, w, h, fill="soft", line="line", radius=False):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLORS[fill]
    shape.line.color.rgb = COLORS[line]
    shape.line.width = Pt(1)
    return shape


def add_bullets(slide, items, x, y, w, h, size=17):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.03)
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = COLORS["ink"]
        p.space_after = Pt(8)
        p.level = 0
    return box


def add_label(slide, text, x, y, w, h, fill="teal_soft", color="teal_dark"):
    shape = add_rect(slide, x, y, w, h, fill=fill, line="line", radius=True)
    tf = shape.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLORS[color]
    return shape


def add_screenshot(slide, filename, x, y, w, h):
    path = ASSETS / filename
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def add_table(slide, rows, x, y, w, h, widths=None, font_size=12):
    table_shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    table = table_shape.table
    if widths:
        for i, width in enumerate(widths):
            table.columns[i].width = Inches(width)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = val
            cell.margin_left = Inches(0.07)
            cell.margin_right = Inches(0.07)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.04)
            fill = COLORS["teal"] if r == 0 else (COLORS["soft"] if r % 2 == 0 else COLORS["white"])
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT
                p.font.size = Pt(font_size)
                p.font.color.rgb = COLORS["white"] if r == 0 else COLORS["ink"]
                p.font.bold = r == 0
    return table_shape


def make_deck():
    prs = Presentation()
    prs.slide_width = Inches(WIDE[0])
    prs.slide_height = Inches(WIDE[1])
    blank = prs.slide_layouts[6]

    slides = []
    for _ in range(12):
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(248, 250, 252)
        slides.append(slide)

    # 1
    s = slides[0]
    add_text(s, "ProjectPilot AI", 0.75, 0.75, 8.5, 0.75, size=42, bold=True)
    add_text(s, "Iteration 1 presentation starts with the live demo", 0.78, 1.55, 8.8, 0.42, size=20, color="muted")
    add_label(s, "Repo version: da43cd6", 0.78, 2.25, 2.55, 0.42)
    add_label(s, "Demo backup: docs/demo-recordings", 3.55, 2.25, 3.25, 0.42)
    add_screenshot(s, "01-new-project.png", 7.25, 1.1, 5.15, 3.6)
    add_bullets(s, [
        "Open http://localhost:3000",
        "Click Run guided workflow",
        "Show requirements, tasks, agents, review, feedback, and final plan",
    ], 0.82, 3.1, 5.6, 1.8, size=20)
    add_footer(s, 1)

    # 2
    s = slides[1]
    add_title(s, "Live Demo Path", "What the audience should see before the slides")
    add_screenshot(s, "02-requirements.png", 6.55, 1.25, 5.75, 4.25)
    add_bullets(s, [
        "The manager agent turns one project idea into requirements",
        "The user approves or edits each requirement before task planning",
        "The system keeps visible checkpoints instead of hidden automation",
    ], 0.8, 1.65, 5.2, 2.1, size=21)
    add_label(s, "Demo point: human approval before the task plan", 0.8, 4.35, 4.7, 0.55)
    add_footer(s, 2)

    # 3
    s = slides[2]
    add_title(s, "What Changed After Feedback", "Main fixes from professor and Team 6 review")
    rows = [
        ["Feedback", "Change in Iteration 1"],
        ["Target user was unclear", "Focused on students building small software projects"],
        ["Workflow was hard to follow", "Added eight-step screen flow and guided demo"],
        ["Agent communication was unclear", "Defined shared context, dependencies, expected output"],
        ["Risk and competitors were missing", "Added risk exposure and competitor comparison"],
        ["Prototype was shown too briefly", "Presentation starts with a live walkthrough"],
    ]
    add_table(s, rows, 0.72, 1.35, 11.9, 4.45, widths=[4.0, 7.9], font_size=15)
    add_footer(s, 3)

    # 4
    s = slides[3]
    add_title(s, "Target User And Problem", "Students need help turning ideas into a clear software plan")
    add_bullets(s, [
        "First users: students building class projects, hackathon apps, or portfolio apps",
        "Problem: project planning requires requirements, task order, dependencies, and review",
        "Common pain: students ask a chatbot for help, then manually coordinate the work",
        "ProjectPilot keeps the plan, task flow, and review checkpoints in one workspace",
    ], 0.8, 1.35, 6.0, 3.25, size=20)
    add_rect(s, 7.35, 1.35, 4.6, 3.25, fill="teal_soft", line="line", radius=True)
    add_text(s, "Scope boundary", 7.65, 1.65, 3.7, 0.3, size=22, bold=True, color="teal_dark")
    add_bullets(s, [
        "Single user",
        "Course-sized software projects",
        "Planning and review first",
        "No full autonomous deployment",
    ], 7.65, 2.1, 3.7, 1.8, size=18)
    add_footer(s, 4)

    # 5
    s = slides[4]
    add_title(s, "Iteration 1 MVP", "The prototype now covers the full planning workflow")
    labels = [
        "Project idea input", "Clarifying questions", "Requirement approval",
        "Task breakdown", "Agent assignment", "Output review",
        "Feedback loop", "Final plan export",
    ]
    x0, y0 = 0.8, 1.35
    for i, label in enumerate(labels):
        x = x0 + (i % 4) * 3.0
        y = y0 + (i // 4) * 1.25
        add_label(s, label, x, y, 2.55, 0.62)
    add_text(s, "Current limitation", 0.82, 4.25, 3.6, 0.35, size=22, bold=True)
    add_bullets(s, [
        "Mock agent responses keep the class demo reliable",
        "The MVP does not call a live AI API yet",
        "The app produces a plan, not deployed code",
    ], 0.82, 4.72, 10.8, 1.2, size=18)
    add_footer(s, 5)

    # 6
    s = slides[5]
    add_title(s, "Task Breakdown Screen", "Manager agent creates dependency-ordered work")
    add_screenshot(s, "03-task-plan.png", 5.95, 1.05, 6.65, 5.55)
    add_bullets(s, [
        "Tasks have an owner agent, dependencies, and expected output",
        "Database and API planning unblock later UI wiring and tests",
        "Users can edit task details before agents run",
        "The plan only runs after approval",
    ], 0.72, 1.45, 4.75, 2.9, size=19)
    add_label(s, "Key improvement: screen flow is now visible", 0.75, 4.75, 4.3, 0.55)
    add_footer(s, 6)

    # 7
    s = slides[6]
    add_title(s, "Agent Communication Contract", "Each specialist receives the same project context")
    rows = [
        ["Agent", "Responsibility", "Returned output"],
        ["Manager", "Requirements, tasks, dependencies, approvals", "Plan updates and notes"],
        ["Frontend", "Screens, UI states, client interactions", "UI plan and edge cases"],
        ["Backend", "Endpoints, validation, retries", "API plan and handler outline"],
        ["Database", "Entities, fields, relationships", "Data model"],
        ["Testing", "Acceptance checks and demo cases", "Test checklist"],
        ["Documentation", "Setup, workflow, privacy notes", "README outline"],
        ["Review", "Missing items and conflicts", "Flags and suggestions"],
    ]
    add_table(s, rows, 0.55, 1.22, 12.25, 4.95, widths=[1.65, 5.25, 5.35], font_size=12)
    add_text(s, "Shared input: task id, description, dependencies, expected output, constraints, and approved project context", 0.65, 6.45, 11.8, 0.32, size=15, color="muted")
    add_footer(s, 7)

    # 8
    s = slides[7]
    add_title(s, "Review And Revision Loop", "The prototype keeps user approval after agent output")
    add_screenshot(s, "04-review-outputs.png", 6.1, 1.05, 6.35, 5.6)
    add_bullets(s, [
        "Review Agent flags missing items and conflicts",
        "Each output needs user approval or revision request",
        "Revised outputs return to pending review",
        "Final plan uses only accepted outputs",
    ], 0.72, 1.4, 4.95, 2.8, size=19)
    add_label(s, "Human oversight is part of the workflow", 0.75, 4.65, 4.45, 0.55)
    add_footer(s, 8)

    # 9
    s = slides[8]
    add_title(s, "Competitor Comparison", "The project stands out by focusing on student workflow and approvals")
    rows = [
        ["Tool", "Strength", "Gap we target"],
        ["ChatGPT", "General help", "User coordinates tasks manually"],
        ["Codex", "Strong coding assistant", "Less focused on project planning flow"],
        ["ChatDev", "Multi-agent software simulation", "Broad and less classroom focused"],
        ["MetaGPT", "Structured agent framework", "Complex for beginners"],
        ["Devika", "Autonomous coding agent", "Less emphasis on approval checkpoints"],
        ["CrewAI", "Agent framework", "Not an end-user student project app"],
    ]
    add_table(s, rows, 0.65, 1.28, 12.0, 4.55, widths=[2.0, 4.0, 6.0], font_size=13)
    add_text(s, "Positioning: a student-facing planning workspace with visible task flow, review, and correction", 0.72, 6.18, 11.4, 0.35, size=17, color="teal_dark", bold=True)
    add_footer(s, 9)

    # 10
    s = slides[9]
    add_title(s, "Risk Exposure", "Highest risk is a wrong task breakdown")
    rows = [
        ["Risk", "Effect", "Exposure", "Mitigation"],
        ["Wrong task breakdown", "20 hrs", "8.0 hrs", "Approve plan before agents run"],
        ["Outputs do not integrate", "18 hrs", "6.3 hrs", "Shared context and Review Agent"],
        ["Unclear requirements", "14 hrs", "4.2 hrs", "Clarifying questions first"],
        ["API integration delay", "16 hrs", "4.0 hrs", "Mock responses until API works"],
        ["Scope too broad", "15 hrs", "4.5 hrs", "Student project MVP only"],
    ]
    add_table(s, rows, 0.55, 1.35, 12.25, 4.45, widths=[3.65, 1.4, 1.55, 5.65], font_size=13)
    add_text(s, "Risk exposure uses RE = probability x effect. Effect and exposure are extra team hours.", 0.65, 6.1, 11.8, 0.35, size=15, color="muted")
    add_footer(s, 10)

    # 11
    s = slides[10]
    add_title(s, "User Feedback Plan", "The next class step is quick feedback on the working prototype")
    add_bullets(s, [
        "Ask 3 to 5 students to use the guided workflow",
        "Record which screen was confusing",
        "Ask whether task approval should happen before agents run",
        "Ask which output would make the system trustworthy",
        "Use feedback to prioritize Iteration 2 UI and API work",
    ], 0.82, 1.35, 6.1, 3.25, size=20)
    add_rect(s, 7.45, 1.35, 4.3, 3.3, fill="soft", line="line", radius=True)
    add_text(s, "Iteration 2 target", 7.75, 1.65, 3.6, 0.32, size=22, bold=True)
    add_bullets(s, [
        "Connect real AI API",
        "Save projects after refresh",
        "Create GitHub issues from tasks",
        "Improve visual agent board",
    ], 7.75, 2.1, 3.4, 1.75, size=18)
    add_footer(s, 11)

    # 12
    s = slides[11]
    add_title(s, "Submission Details", "Files are in GitHub with versioned code and demo backup")
    add_bullets(s, [
        "Repository: github.com/manit0700/projectpilot-ai-cse3311",
        "Version for presentation: da43cd6",
        "Demo backup video: docs/demo-recordings/projectpilot-ai-demo.webm",
        "Written instructions: docs/GITHUB_TEAM_INSTRUCTIONS.md",
    ], 0.82, 1.35, 10.5, 1.8, size=20)
    add_text(s, "Sources", 0.82, 3.75, 2.0, 0.35, size=22, bold=True)
    add_bullets(s, [
        "Professor lecture guidance from CSE 3311 meeting transcript, September 17, 2026",
        "Team 6 inception review action items",
        "Competitor references: ChatGPT, Codex, ChatDev, MetaGPT, Devika, CrewAI public project/product pages",
    ], 0.82, 4.2, 10.9, 1.45, size=16)
    add_label(s, "Q&A", 0.82, 6.05, 1.4, 0.55)
    add_footer(s, 12)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    OUT.parent.mkdir(parents=True, exist_ok=True)
    make_deck()
